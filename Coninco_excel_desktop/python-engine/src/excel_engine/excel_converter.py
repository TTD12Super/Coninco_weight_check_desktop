import pandas as pd

from .quantity_calculator import calculate_K
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelConverter:
    def __init__(self, dataFrame: pd.DataFrame):
        self.dataFrame = dataFrame.copy()

    def remove_columns(self, columns: list[str]) -> None:
        self.dataFrame = self.dataFrame.drop(
            columns=columns,
            errors="ignore"
        )

    def remove_empty_rows(self, columns: list[str]) -> None:
        self.dataFrame = self.dataFrame.dropna(
            subset=columns,
            how="all"
        )

    def remove_rows_with_null(self, columns: list[str]) -> None:
        self.dataFrame = self.dataFrame.dropna(
            subset=columns,
            how="any"
        )

    def remove_duplicate_rows(self) -> None:
        self.dataFrame = self.dataFrame.drop_duplicates()

    def rename_columns(self, mapping: dict[str, str]) -> None:
        self.dataFrame = self.dataFrame.rename(columns=mapping)

    def set_column(self, column: str, values) -> None:
        self.dataFrame[column] = values

    def add_k_by_function(self) -> None:
        """Add K calculated from plan, up/down, and vertical-axis cells."""
        self.dataFrame["K_By_function"] = self.dataFrame.apply(
            calculate_K,
            axis=1,
        )

    def get_data(self) -> pd.DataFrame:
        return self.dataFrame

    def get_data(self) -> pd.DataFrame:
        return self.dataFrame

    def preview(self, rows: int = 10) -> pd.DataFrame:
        return self.dataFrame.head(rows)
    
    def export_excel(self, output_path: str) -> None:
        self.dataFrame.to_excel(
            output_path,
            index=False,
            engine="openpyxl"
        )

        from openpyxl import load_workbook

        workbook = load_workbook(output_path)
        worksheet = workbook.active

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True
                )

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[column_letter].width = min(
                max_length + 2,
                40
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        workbook.save(output_path)
