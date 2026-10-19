"""Calculate contractor rows without removing the original workbook layout."""

from __future__ import annotations

from copy import copy
import csv
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries

from .excel_reader import ExcelReader
from .product_formatter import ProductFormat, calculate_section
from .quantity_calculator import calculate_K


# Original contractor sheet positions, zero-based.
_SOURCE_COLUMNS = {
    "item_no": 1,
    "description": 2,
    "unit": 3,
    "unit_duplicate": 4,
    "W": 5,
    "H": 6,
    "W1": 7,
    "H1": 8,
    "W2": 9,
    "H2": 10,
    "L_M": 11,
    "R_D": 12,
    "a": 13,
    "t": 14,
    "Q": 15,
    "K": 16,
    "plan": 17,
    "up_down": 18,
    "vertical_axis": 19,
    "formula_or_insulation_detail": 22,
    "acoustic_fireproof_detail": 23,
    "note": 24,
}

_K_COLUMN = 16
_OUTER_S_COLUMN = 20
_INNER_S_COLUMN = 21

_K_OUTPUT = "K (tính bằng máy)"
_OUTER_S_OUTPUT = "S (cách nhiệt tính bằng máy)"
_INNER_S_OUTPUT = "S (cách nhiệt trong tính bằng máy)"

_DISPLAY_HEADERS = {
    _K_OUTPUT: "K\n(tính bằng máy)",
    _OUTER_S_OUTPUT: "S\n(cách nhiệt tính bằng máy)",
    _INNER_S_OUTPUT: "S\n(cách nhiệt trong tính bằng máy)",
}


def read_raw_workbook(path: str | Path) -> dict[str, pd.DataFrame]:
    """Read every used cell, including the original title and header rows."""
    reader = ExcelReader(path)
    return pd.read_excel(
        reader.path,
        sheet_name=None,
        engine=reader.engine,
        header=None,
        skiprows=0,
    )


def _cell(row: pd.Series, name: str) -> Any:
    return row.iloc[_SOURCE_COLUMNS[name]] if _SOURCE_COLUMNS[name] < len(row) else None


def _normalized_row(row: pd.Series) -> dict[str, Any]:
    normalized = {name: _cell(row, name) for name in _SOURCE_COLUMNS}
    normalized["section"] = calculate_section(normalized)
    return normalized


def _has_value(value: Any) -> bool:
    return value is not None and not (not isinstance(value, str) and pd.isna(value)) and str(value).strip() != ""


def _is_numeric(value: Any) -> bool:
    if not _has_value(value):
        return False
    try:
        float(str(value).replace(",", ".").strip())
    except (TypeError, ValueError):
        return False
    return True


def _is_calculable(row: dict[str, Any]) -> bool:
    return _has_value(row["description"]) and any(
        _is_numeric(row[name])
        for name in ("W", "H", "W1", "H1", "W2", "H2", "L_M", "R_D")
    )


def _number(value: Any) -> float:
    """Convert an Excel number to a comparable value; blank cells equal zero."""
    if not _has_value(value):
        return 0.0
    try:
        number = float(str(value).replace(",", ".").strip())
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if math.isnan(number) else number


def _text_value(value: Any) -> str:
    """Return a display-safe workbook value without exposing pandas NaN."""
    return "" if not _has_value(value) else str(value).replace("\n", " ").replace("\r", " ")


def _matches_original(original: Any, calculated: float) -> bool:
    """Compare displayed three-decimal values, allowing normal rounding noise."""
    return abs(_number(original) - calculated) <= 0.0005


def calculate_raw_sheet(
    data: pd.DataFrame,
    section: str | None = None,
    product_type: str | None = None,
) -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]], list[dict[str, Any]]]:
    """Calculate rows while retaining every original row in ``data``."""
    calculated: list[dict[str, Any]] = []
    output_values: dict[int, dict[str, Any]] = {}
    errors: list[dict[str, Any]] = []

    for index, raw_row in data.iterrows():
        row = _normalized_row(raw_row)
        if not _is_calculable(row):
            continue

        source_row = int(index) + 1
        try:
            row["K_By_function"] = calculate_K(row)
            product = ProductFormat.from_row(
                row,
                section=section,
                product_type=product_type,
            )
            result = product.calculate(row)
            inputs = product.inputs_from(row)
            calculated_row = {
                "source_row": source_row,
                "description": product.description,
                "section": product.section,
                "product_type": product.product_type,
                "inputs": inputs.__dict__,
                "result": result.to_dict(),
            }
            calculated.append(calculated_row)
            output_values[index] = {
                "K": inputs.K,
                "S": result.R,
                "S_inner": result.T,
                "matches": {
                    "K": _matches_original(row["K"], inputs.K),
                    "S": _matches_original(raw_row.iloc[_OUTER_S_COLUMN], result.R),
                    "S_inner": _matches_original(raw_row.iloc[_INNER_S_COLUMN], result.T),
                },
            }
        except (TypeError, ValueError) as error:
            errors.append(
                {
                    "source_row": source_row,
                    "description": row.get("description", ""),
                    "message": str(error),
                }
            )

    return calculated, output_values, errors


def _header_row(data: pd.DataFrame) -> int:
    for index, row in data.iterrows():
        text = " ".join(str(value).lower() for value in row if _has_value(value))
        if "total of quantities" in text and "khối lượng cách nhiệt" in text:
            return int(index)
    return 0


def _insert_result_columns(
    data: pd.DataFrame,
    output_values: dict[int, dict[str, Any]],
) -> pd.DataFrame:
    """Insert calculated columns beside original K, S, and inner-S columns."""
    result = data.copy()
    header = _header_row(data)

    def values(key: str, label: str) -> list[Any]:
        output = [None] * len(data)
        output[header] = label
        for index, row_values in output_values.items():
            output[index] = row_values[key]
        return output

    # Insert from left to right using positions after previous insertions.
    result.insert(_K_COLUMN + 1, _K_OUTPUT, values("K", _K_OUTPUT))
    result.insert(_OUTER_S_COLUMN + 2, _OUTER_S_OUTPUT, values("S", _OUTER_S_OUTPUT))
    result.insert(_INNER_S_COLUMN + 3, _INNER_S_OUTPUT, values("S_inner", _INNER_S_OUTPUT))
    return result


def _soffice_path() -> str:
    """Find LibreOffice, used to convert legacy .xls files without losing layout."""
    executable = shutil.which("soffice")
    if executable:
        return executable

    bundled = (
        Path.home()
        / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"
    )
    if bundled.is_file():
        return str(bundled)
    raise FileNotFoundError(
        "Không tìm thấy LibreOffice (soffice) để chuyển file .xls sang .xlsx."
    )


def _xlsx_source(source: Path, temporary_directory: Path) -> Path:
    """Return an xlsx source while keeping the original .xls file untouched."""
    if source.suffix.lower() == ".xlsx":
        return source

    subprocess.run(
        [
            _soffice_path(),
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(temporary_directory),
            str(source),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    converted = temporary_directory / f"{source.stem}.xlsx"
    if not converted.is_file():
        raise FileNotFoundError(f"Không tạo được file xlsx tạm từ {source}.")
    return converted


def _shift_merged_ranges(worksheet, insert_at: int) -> None:
    """Move merge ranges around an inserted column.

    openpyxl moves cell values for ``insert_cols`` but does not update merged
    ranges.  Rebuilding the ranges here keeps the original header layout.
    """
    ranges = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
    for cell_range in ranges:
        worksheet.unmerge_cells(cell_range)

    for cell_range in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        if min_col >= insert_at:
            min_col += 1
            max_col += 1
        elif max_col >= insert_at:
            max_col += 1
        worksheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )


def _copy_column_dimension(worksheet, source_column: int, target_column: int) -> None:
    source = worksheet.column_dimensions[get_column_letter(source_column)]
    target = worksheet.column_dimensions[get_column_letter(target_column)]
    target.width = source.width
    target.hidden = source.hidden
    target.bestFit = source.bestFit
    target.outlineLevel = source.outlineLevel
    target.collapsed = source.collapsed


def _copy_column_style(
    worksheet,
    source_column: int,
    target_column: int,
    red_text: bool = True,
    width: float | None = None,
) -> None:
    """Copy the original column's cell styles to a newly inserted column."""
    _copy_column_dimension(worksheet, source_column, target_column)
    if width is not None:
        worksheet.column_dimensions[get_column_letter(target_column)].width = width
    for row in range(1, worksheet.max_row + 1):
        source = worksheet.cell(row=row, column=source_column)
        target = worksheet.cell(row=row, column=target_column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        if red_text:
            font = copy(target.font)
            font.color = "FFFF0000"
            target.font = font


def _set_new_header(worksheet, column: int, row: int, label: str) -> None:
    """Write a new header without breaking an existing merged cell."""
    cell = worksheet.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        for candidate_row in range(row + 1, worksheet.max_row + 1):
            candidate = worksheet.cell(row=candidate_row, column=column)
            if not isinstance(candidate, MergedCell):
                cell = candidate
                break
    cell.value = label
    font = copy(cell.font)
    font.bold = True
    font.color = "FFFF0000"
    cell.font = font
    alignment = copy(cell.alignment)
    alignment.horizontal = "center"
    alignment.vertical = "center"
    alignment.wrap_text = True
    cell.alignment = alignment
    current_height = worksheet.row_dimensions[row].height or 0
    worksheet.row_dimensions[row].height = max(current_height, 42)


def _write_preview(
    destination: Path,
    data: pd.DataFrame,
    calculated: list[dict[str, Any]],
    output_values: dict[int, dict[str, Any]],
) -> Path:
    """Write a compact sidecar for the desktop table preview."""
    preview = destination.with_name(f"{destination.stem}_preview.tsv")
    headers = [
        "source_row", "item_no", "description", "unit", "unit_duplicate",
        "section", "product_type",
        "w", "h", "w1", "h1", "w2", "h2", "l_m", "r_d", "a", "t", "q",
        "plan", "up_down", "vertical_axis", "formula_detail", "acoustic_fireproof_detail", "note",
        "k_by_function", "contractor_k", "calculated_k", "k_match",
        "contractor_s", "calculated_s", "s_match", "contractor_inner_s",
        "calculated_inner_s", "inner_match",
    ]
    with preview.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
        writer.writerow(headers)
        for item in calculated:
            source_row = int(item["source_row"])
            raw_row = data.iloc[source_row - 1]
            values = output_values[source_row - 1]
            result = item["result"]["quantities"]
            inputs = item["inputs"]
            writer.writerow([
                source_row,
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["item_no"]]),
                _text_value(item["description"]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["unit"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["unit_duplicate"]]),
                item["section"],
                item["product_type"],
                inputs["w"],
                inputs["h"],
                inputs["w1"],
                inputs["h1"],
                inputs["w2"],
                inputs["h2"],
                inputs["L_M"],
                inputs["R_D"],
                inputs["a"],
                inputs["t"],
                inputs["Q"],
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["plan"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["up_down"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["vertical_axis"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["formula_or_insulation_detail"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["acoustic_fireproof_detail"]]),
                _text_value(raw_row.iloc[_SOURCE_COLUMNS["note"]]),
                inputs["K"],
                _number(raw_row.iloc[_K_COLUMN]),
                inputs["K"],
                values["matches"]["K"],
                _number(raw_row.iloc[_OUTER_S_COLUMN]),
                result["R"],
                values["matches"]["S"],
                _number(raw_row.iloc[_INNER_S_COLUMN]),
                result["T"],
                values["matches"]["S_inner"],
            ])
    return preview


def export_calculated_workbook(
    input_path: str | Path,
    output_path: str | Path | None = None,
    sheet_name: str = "02",
) -> dict[str, Any]:
    """Export a full workbook while preserving the source workbook layout."""
    source = Path(input_path).expanduser().resolve()
    destination = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else source.with_name(f"{source.stem}_calculated.xlsx")
    )
    if source == destination:
        raise ValueError("The output workbook must be different from the input workbook.")

    if destination.suffix.lower() != ".xlsx":
        raise ValueError("File kết quả phải có đuôi .xlsx để giữ nguyên định dạng.")

    workbook = read_raw_workbook(source)
    if sheet_name not in workbook:
        raise ValueError(f"Worksheet not found: {sheet_name}")

    calculated, output_values, errors = calculate_raw_sheet(workbook[sheet_name])
    mismatch_count = sum(
        not match
        for values in output_values.values()
        for match in values["matches"].values()
    )

    # Read and edit the actual workbook instead of recreating it from a DataFrame.
    # This preserves sheet names, merged cells, row heights, fonts and cell styles.
    with tempfile.TemporaryDirectory(prefix="excel_auditor_") as temporary_name:
        xlsx_source = _xlsx_source(source, Path(temporary_name))
        workbook_out = load_workbook(xlsx_source)
        worksheet = workbook_out[sheet_name]

        # Insert from left to right. These are Excel's 1-based column numbers.
        insertions = (
            (18, 17, _K_OUTPUT, "K", 18),
            (23, 22, _OUTER_S_OUTPUT, "S", 28),
            (25, 24, _INNER_S_OUTPUT, "S_inner", 34),
        )
        for insert_at, source_column, label, value_key, width in insertions:
            _shift_merged_ranges(worksheet, insert_at)
            worksheet.insert_cols(insert_at, 1)
            _copy_column_style(worksheet, source_column, insert_at, width=width)
            _set_new_header(
                worksheet,
                insert_at,
                _header_row(workbook[sheet_name]) + 1,
                _DISPLAY_HEADERS[label],
            )
            for raw_index, values in output_values.items():
                cell = worksheet.cell(
                    row=int(raw_index) + 1,
                    column=insert_at,
                )
                cell.value = values[value_key]
                font = copy(cell.font)
                font.color = (
                    "FF008000"
                    if values["matches"][value_key]
                    else "FFFF0000"
                )
                cell.font = font

        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook_out.save(destination)

    preview_path = _write_preview(destination, workbook[sheet_name], calculated, output_values)

    return {
        "input": str(source),
        "output": str(destination),
        "sheet": sheet_name,
        "row_count": len(workbook[sheet_name]),
        "calculated_count": len(calculated),
        "mismatch_count": mismatch_count,
        "preview": str(preview_path),
        "errors": errors,
    }


__all__ = ["calculate_raw_sheet", "export_calculated_workbook", "read_raw_workbook"]
